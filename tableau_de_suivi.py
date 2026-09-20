import datetime
import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import streamlit as st

st.set_page_config(
    page_title="Suivi Non-Conformités Cuisine",
    page_icon="🍳",
    layout="wide",
)

st.title("🍳 Plateforme de Suivi des Non-Conformités - Cuisine Centrale")
st.write(
    "Saisissez vos contrôles HACCP, gérez vos données en direct et exportez le rapport Excel."
)

# Initialisation de la base de données dans la session Streamlit
if "df_nc" not in st.session_state:
    st.session_state.df_nc = pd.DataFrame(
        [
            {
                "ID": "NC-2026-001",
                "Date Détection": "2026-03-15",
                "Secteur / Zone": "Chambre Froide Positive 1",
                "Catégorie HACCP": "Chaîne du froid",
                "Description de la Non-Conformité": "Température relevée à +8°C au lieu de +3°C max.",
                "Niveau de Risque": "Critique",
                "Action Corrective Immédiate": "Mise en isolement des denrées et transfert vers CF2.",
                "Responsable Action": "Chef Équipe Froid",
                "Date Échéance": "2026-03-15",
                "Statut": "Clôturé",
            },
            {
                "ID": "NC-2026-002",
                "Date Détection": "2026-03-16",
                "Secteur / Zone": "Zone Réception / Quai",
                "Catégorie HACCP": "Traçabilité & Livraison",
                "Description de la Non-Conformité": "Colis de poisson frais reçus sans étiquette de traçabilité.",
                "Niveau de Risque": "Majeur",
                "Action Corrective Immédiate": "Refus de la livraison auprès du fournisseur.",
                "Responsable Action": "Responsable Achats",
                "Date Échéance": "2026-03-16",
                "Statut": "Clôturé",
            },
            {
                "ID": "NC-2026-003",
                "Date Détection": "2026-03-18",
                "Secteur / Zone": "Légumerie / Machines à glaçons",
                "Catégorie HACCP": "Hygiène & Nettoyage",
                "Description de la Non-Conformité": "Machine à glaçons : absence de contrôle de propreté des surfaces en contact avec la glace.",
                "Niveau de Risque": "Majeur",
                "Action Corrective Immédiate": "Nettoyage complet et vérification des joints et bacs.",
                "Responsable Action": "Responsable Stewarding",
                "Date Échéance": "2026-03-20",
                "Statut": "En cours",
            },
        ]
    )

# ---------------------------------------------------------
# BARRE LATÉRALE : FORMULAIRE DE SAISIE
# ---------------------------------------------------------
st.sidebar.header("➕ Saisir une Non-Conformité")

with st.sidebar.form(key="form_saisie", clear_on_submit=True):
    next_id = f"NC-2026-00{len(st.session_state.df_nc) + 1}"
    date_detec = st.date_input(
        "Date de Détection", value=datetime.date.today()
    )

    secteur = st.selectbox(
        "Secteur / Zone",
        [
            "Chambre Froide Positive",
            "Chambre Froide Négative",
            "Zone Réception / Quai",
            "Cuisine Chaude",
            "Pâtisserie",
            "Légumerie",
            "Zone Plonge / Stewarding",
            "Local Glaçons",
        ],
    )

    categorie = st.selectbox(
        "Catégorie HACCP",
        [
            "Chaîne du froid",
            "Traçabilité & Livraison",
            "Hygiène & Nettoyage",
            "Matériel & Maintenance",
            "Tenue & Hygiène Personnel",
            "Gestion des Déchets",
        ],
    )

    description = st.text_area("Description de la Non-Conformité")
    risque = st.select_slider(
        "Niveau de Risque", options=["Mineur", "Majeur", "Critique"]
    )
    action = st.text_area("Action Corrective Immédiate")
    responsable = st.text_input("Responsable de l'Action")
    date_echeance = st.date_input("Date Échéance", value=datetime.date.today())
    statut = st.selectbox("Statut", ["Ouvert", "En cours", "Clôturé"])

    btn_submit = st.form_submit_button(label="Enregistrer la NC")

if btn_submit:
    nouvelle_nc = {
        "ID": next_id,
        "Date Détection": str(date_detec),
        "Secteur / Zone": secteur,
        "Catégorie HACCP": categorie,
        "Description de la Non-Conformité": description,
        "Niveau de Risque": risque,
        "Action Corrective Immédiate": action,
        "Responsable Action": responsable,
        "Date Échéance": str(date_echeance),
        "Statut": statut,
    }

    # Ajout au dataframe de la session
    st.session_state.df_nc = pd.concat(
        [st.session_state.df_nc, pd.DataFrame([nouvelle_nc])], ignore_index=True
    )
    st.sidebar.success(f"Non-conformité {next_id} enregistrée !")

# ---------------------------------------------------------
# CORPS PRINCIPAL : AFFICHAGE ET ÉDITION DU TABLEAU
# ---------------------------------------------------------
st.subheader("📋 Tableau dynamique des Non-Conformités")
st.caption(
    "Vous pouvez modifier directement les cases dans le tableau ci-dessous ou ajouter/supprimer des lignes."
)

# Tableau interactif (st.data_editor)
edited_df = st.data_editor(
    st.session_state.df_nc,
    num_rows="dynamic",
    use_container_width=True,
    column_config={
        "Statut": st.column_config.SelectboxColumn(
            "Statut", options=["Ouvert", "En cours", "Clôturé"], required=True
        ),
        "Niveau de Risque": st.column_config.SelectboxColumn(
            "Niveau de Risque",
            options=["Mineur", "Majeur", "Critique"],
            required=True,
        ),
    },
)

# Mettre à jour la session avec les modifications directes
st.session_state.df_nc = edited_df


# ---------------------------------------------------------
# GÉNÉRATION DU FICHIER EXCEL
# ---------------------------------------------------------
def generer_excel(df_data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Suivi NC Cuisine"

    # Titre
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "SUIVI DES NON-CONFORMITÉS - CUISINE CENTRALE HÔTEL"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # En-têtes
    headers = list(df_data.columns)
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.value = header_title
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="2F5597", end_color="2F5597", fill_type="solid"
        )
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    ws.row_dimensions[3].height = 25

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for row_idx, row_data in enumerate(df_data.values, 4):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.font = Font(name="Arial", size=9)
            cell.border = thin_border
            if col_idx in [1, 2, 6, 9, 10]:
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )
            else:
                cell.alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )

    # Mise en forme conditionnelle des Statuts
    red_fill = PatternFill(
        start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
    )
    red_font = Font(color="C00000", bold=True, name="Arial", size=9)
    yellow_fill = PatternFill(
        start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
    )
    yellow_font = Font(color="B25900", bold=True, name="Arial", size=9)
    green_fill = PatternFill(
        start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
    )
    green_font = Font(color="375623", bold=True, name="Arial", size=9)

    last_row = len(df_data) + 3
    statut_range = f"J4:J{last_row}"

    ws.conditional_formatting.add(
        statut_range,
        CellIsRule(
            operator="equal",
            formula=['"Ouvert"'],
            fill=red_fill,
            font=red_font,
        ),
    )
    ws.conditional_formatting.add(
        statut_range,
        CellIsRule(
            operator="equal",
            formula=['"En cours"'],
            fill=yellow_fill,
            font=yellow_font,
        ),
    )
    ws.conditional_formatting.add(
        statut_range,
        CellIsRule(
            operator="equal",
            formula=['"Clôturé"'],
            fill=green_fill,
            font=green_font,
        ),
    )

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


st.subheader("📥 Exporter le rapport Excel")
excel_data = generer_excel(st.session_state.df_nc)

st.download_button(
    label="Télécharger le fichier Excel mis à jour (.xlsx)",
    data=excel_data,
    file_name="Suivi_Non_Conformites_Cuisine.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
